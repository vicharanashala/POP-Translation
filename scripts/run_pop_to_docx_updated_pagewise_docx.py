import argparse
import json
import os
import re
import shutil
import subprocess
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

import fitz  # PyMuPDF
from bs4 import BeautifulSoup, NavigableString
from google import genai
from google.genai import types
from docx import Document
from docxcompose.composer import Composer


# ---------------------------------------------------------------------
# Runtime logging
# ---------------------------------------------------------------------

LOG_FILE_PATH: Path | None = None
LOG_LOCK = threading.Lock()


def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def format_duration(seconds: float) -> str:
    seconds = int(seconds)

    hours = seconds // 3600
    minutes = (seconds % 3600) // 60
    secs = seconds % 60

    if hours:
        return f"{hours}h {minutes}m {secs}s"
    if minutes:
        return f"{minutes}m {secs}s"
    return f"{secs}s"


def set_log_file(log_file_path: Path):
    global LOG_FILE_PATH
    LOG_FILE_PATH = log_file_path
    LOG_FILE_PATH.parent.mkdir(parents=True, exist_ok=True)

    with LOG_FILE_PATH.open("a", encoding="utf-8") as f:
        f.write("\n" + "=" * 100 + "\n")
        f.write(f"Pipeline run started at {now_str()}\n")
        f.write("=" * 100 + "\n")


def log(message: str):
    line = f"[{now_str()}] {message}"

    with LOG_LOCK:
        print(line, flush=True)

        if LOG_FILE_PATH:
            with LOG_FILE_PATH.open("a", encoding="utf-8") as f:
                f.write(line + "\n")


# ---------------------------------------------------------------------
# Prompt fallback: used only if --prompt-file is missing/not supplied.
# Prefer using prompts/page_to_pdf.txt from your repo.
# ---------------------------------------------------------------------

DEFAULT_PROMPT = """Translate this PDF page into English and return only clean HTML. Follow these rules strictly: 1. Preserve the original page structure as faithfully as possible. 2. Preserve headings, paragraphs, numbered lists, bullet points, tables, captions, and reading order. 3. Preserve tables strictly in valid HTML table format. 4. Do not convert tables into paragraphs, bullet points, or free text. 5. Preserve the original number of rows and columns as closely as possible. 6. Preserve chemical names, crop names, pest names, formulation codes, units, doses, percentages, and numbers exactly. 7. Do not omit any visible textual content. 8. Do not summarize, paraphrase, simplify, or explain. 9. Do not add content that is not present in the PDF. 10. Return only HTML. Do not return Markdown. 11. Use semantic HTML tags where appropriate, such as h1-h6, p, ul, ol, li, table, thead, tbody, tr, th, td, figure, figcaption, div, and span. 12. Preserve multi-column or visually separated sections as separate HTML blocks in reading order. 13. Do not invent image descriptions. If the page contains an image or figure whose original binary content cannot be reproduced in HTML output, preserve its position using a minimal placeholder block such as: [IMAGE] 14. If the source page already contains a visible caption, preserve that caption near the corresponding image placeholder. 15. Do not use code fences.
Return a single self-contained HTML fragment for this one PDF page."""


# ---------------------------------------------------------------------
# General utilities
# ---------------------------------------------------------------------

def load_prompt(prompt_path: Path | None) -> str:
    if prompt_path and prompt_path.exists():
        log(f"Prompt loaded from: {prompt_path}")
        return prompt_path.read_text(encoding="utf-8").strip()

    log("[WARN] Prompt file not found. Using fallback DEFAULT_PROMPT.")
    return DEFAULT_PROMPT.strip()


def clean_html(text: str) -> str:
    """
    Same cleaning behavior as notebook 02:
    remove accidental markdown fences and return final HTML with trailing newline.
    """
    text = text.strip()

    text = re.sub(r"^```html\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"^```\s*", "", text)
    text = re.sub(r"\s*```$", "", text)

    return text.strip() + "\n"


def page_name(page_num: int) -> str:
    return f"page_{page_num:03d}"


def safe_docx_name(text: str) -> str:
    """
    Keep readable filename but remove Windows-problematic characters.
    """
    text = re.sub(r'[<>:"/\\|?*]', "_", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


# ---------------------------------------------------------------------
# Step 1: Split source PDF into page-wise PDFs
# Based on notebook 01.
# ---------------------------------------------------------------------

def split_pdf_to_page_folders(
    source_pdf_path: Path,
    output_root: Path,
    start_page: int,
    end_page: int | None,
    overwrite_existing: bool = False,
):
    stage_start = time.perf_counter()

    if not source_pdf_path.exists():
        raise FileNotFoundError(f"Source PDF not found: {source_pdf_path}")

    doc = fitz.open(source_pdf_path)
    total_pages = len(doc)

    if end_page is None:
        end_page = total_pages

    if start_page < 1:
        raise ValueError("start_page must be >= 1")

    if end_page > total_pages:
        raise ValueError(f"end_page cannot be greater than total PDF pages: {total_pages}")

    if start_page > end_page:
        raise ValueError("start_page cannot be greater than end_page")

    output_root.mkdir(parents=True, exist_ok=True)
    saved_files = []
    created_count = 0
    skipped_count = 0

    for i in range(start_page - 1, end_page):
        pnum = i + 1
        pname = page_name(pnum)
        page_dir = output_root / pname
        page_dir.mkdir(parents=True, exist_ok=True)

        out_pdf_path = page_dir / f"{pname}.pdf"

        if out_pdf_path.exists() and not overwrite_existing:
            saved_files.append(out_pdf_path)
            skipped_count += 1
            continue

        new_doc = fitz.open()
        new_doc.insert_pdf(doc, from_page=i, to_page=i)
        new_doc.save(out_pdf_path)
        new_doc.close()

        saved_files.append(out_pdf_path)
        created_count += 1

    doc.close()

    elapsed = time.perf_counter() - stage_start
    log(
        f"PDF split complete | total_pdf_pages={total_pages} | "
        f"selected_pages={len(saved_files)} | created={created_count} | "
        f"reused={skipped_count} | duration={format_duration(elapsed)}"
    )

    return saved_files, total_pages


# ---------------------------------------------------------------------
# Step 2: Translate page PDF to HTML
# This intentionally follows notebook 02:
# - gemini-3.1-pro-preview default
# - generate_content_stream
# - direct bytes PDF input
# - thinking_level="HIGH"
# - googleSearch enabled
# ---------------------------------------------------------------------

def make_generate_content_config(enable_google_search: bool = True):
    tools = []

    if enable_google_search:
        tools = [
            types.Tool(
                googleSearch=types.GoogleSearch()
            ),
        ]

    return types.GenerateContentConfig(
        thinking_config=types.ThinkingConfig(
            thinking_level="HIGH",
        ),
        tools=tools,
    )


def translate_page_pdf_to_html(
    pdf_path: Path,
    prompt: str,
    model: str,
    api_key: str,
    max_retries: int,
    retry_wait_seconds: int,
    enable_google_search: bool,
) -> str:
    pdf_bytes = pdf_path.read_bytes()
    last_error = None

    client = genai.Client(api_key=api_key)
    generate_content_config = make_generate_content_config(
        enable_google_search=enable_google_search
    )

    pname = pdf_path.parent.name
    pdf_size_kb = len(pdf_bytes) / 1024

    for attempt in range(1, max_retries + 1):
        attempt_start = time.perf_counter()

        try:
            collected = []

            log(
                f"{pname} | attempt {attempt}/{max_retries} | "
                f"request sent to Gemini | pdf_size={pdf_size_kb:.1f} KB"
            )

            chunk_count = 0
            text_chunk_count = 0
            last_stream_log = time.perf_counter()

            for chunk in client.models.generate_content_stream(
                model=model,
                contents=[
                    types.Content(
                        role="user",
                        parts=[
                            types.Part.from_text(text=prompt),
                            types.Part.from_bytes(
                                data=pdf_bytes,
                                mime_type="application/pdf",
                            ),
                        ],
                    ),
                ],
                config=generate_content_config,
            ):
                chunk_count += 1

                if text := chunk.text:
                    collected.append(text)
                    text_chunk_count += 1

                now = time.perf_counter()
                if now - last_stream_log >= 30:
                    log(
                        f"{pname} | attempt {attempt} | still streaming... "
                        f"elapsed={format_duration(now - attempt_start)} | "
                        f"chunks={chunk_count} | text_chunks={text_chunk_count}"
                    )
                    last_stream_log = now

            full_text = "".join(collected).strip()
            attempt_elapsed = time.perf_counter() - attempt_start

            if not full_text:
                raise RuntimeError("Empty response received from Gemini.")

            log(
                f"{pname} | attempt {attempt} | Gemini response completed | "
                f"duration={format_duration(attempt_elapsed)} | "
                f"chunks={chunk_count} | text_chunks={text_chunk_count} | chars={len(full_text)}"
            )

            return clean_html(full_text)

        except Exception as e:
            attempt_elapsed = time.perf_counter() - attempt_start
            last_error = e

            log(
                f"{pname} | attempt {attempt}/{max_retries} failed | "
                f"duration={format_duration(attempt_elapsed)} | error={e}"
            )

            if attempt < max_retries:
                log(f"{pname} | waiting {retry_wait_seconds}s before retry")
                time.sleep(retry_wait_seconds)

    raise RuntimeError(f"Failed after {max_retries} attempts for {pdf_path}") from last_error


def process_translation_for_page(
    page_pdf_path: Path,
    prompt: str,
    model: str,
    api_key: str,
    overwrite_existing: bool,
    max_retries: int,
    retry_wait_seconds: int,
    enable_google_search: bool,
):
    page_start = time.perf_counter()

    page_dir = page_pdf_path.parent
    pname = page_dir.name
    html_path = page_dir / "translated.html"
    err_path = page_dir / "translated_error.txt"

    if html_path.exists() and not overwrite_existing:
        log(f"{pname} | translation skipped | translated.html already exists")
        return {
            "page": pname,
            "status": "skipped",
            "reason": "translated.html already exists",
            "output_file": str(html_path),
            "duration_seconds": 0,
            "started_at": now_str(),
            "finished_at": now_str(),
        }

    log(f"{pname} | translation started")

    try:
        translated_html = translate_page_pdf_to_html(
            pdf_path=page_pdf_path,
            prompt=prompt,
            model=model,
            api_key=api_key,
            max_retries=max_retries,
            retry_wait_seconds=retry_wait_seconds,
            enable_google_search=enable_google_search,
        )

        html_path.write_text(translated_html, encoding="utf-8")

        page_elapsed = time.perf_counter() - page_start

        log(
            f"{pname} | translation finished | "
            f"duration={format_duration(page_elapsed)} | saved={html_path}"
        )

        return {
            "page": pname,
            "status": "success",
            "output_file": str(html_path),
            "duration_seconds": round(page_elapsed, 2),
            "started_at": None,
            "finished_at": now_str(),
        }

    except Exception as e:
        page_elapsed = time.perf_counter() - page_start

        err = str(e)
        err_path.write_text(err, encoding="utf-8")

        log(
            f"{pname} | translation failed | "
            f"duration={format_duration(page_elapsed)} | error={err}"
        )

        return {
            "page": pname,
            "status": "failed",
            "error": err,
            "error_file": str(err_path),
            "duration_seconds": round(page_elapsed, 2),
            "started_at": None,
            "finished_at": now_str(),
        }


def translate_pages(
    page_pdf_files,
    prompt: str,
    model: str,
    api_key: str,
    overwrite_existing: bool,
    max_retries: int,
    retry_wait_seconds: int,
    enable_google_search: bool,
    concurrency: int,
):
    """
    concurrency=1 gives notebook-equivalent sequential behavior.
    concurrency>1 runs independent page translations in parallel.
    """
    stage_start = time.perf_counter()
    summary = []

    total_pages = len(page_pdf_files)
    log(
        f"Translation stage started | pages={total_pages} | "
        f"concurrency={concurrency} | model={model} | "
        f"google_search={enable_google_search}"
    )

    if concurrency <= 1:
        for idx, pdf_path in enumerate(page_pdf_files, start=1):
            log(f"Translation progress | starting {idx}/{total_pages} | page={pdf_path.parent.name}")
            result = process_translation_for_page(
                page_pdf_path=pdf_path,
                prompt=prompt,
                model=model,
                api_key=api_key,
                overwrite_existing=overwrite_existing,
                max_retries=max_retries,
                retry_wait_seconds=retry_wait_seconds,
                enable_google_search=enable_google_search,
            )
            summary.append(result)
            log(f"Translation progress | completed {idx}/{total_pages} | page={pdf_path.parent.name}")

        elapsed = time.perf_counter() - stage_start
        log_translation_summary(summary, elapsed)
        return summary

    with ThreadPoolExecutor(max_workers=concurrency) as executor:
        future_map = {
            executor.submit(
                process_translation_for_page,
                pdf_path,
                prompt,
                model,
                api_key,
                overwrite_existing,
                max_retries,
                retry_wait_seconds,
                enable_google_search,
            ): pdf_path
            for pdf_path in page_pdf_files
        }

        completed = 0

        for future in as_completed(future_map):
            pdf_path = future_map[future]
            result = future.result()
            summary.append(result)

            completed += 1
            log(
                f"Translation progress | completed {completed}/{total_pages} | "
                f"page={pdf_path.parent.name} | status={result.get('status')}"
            )

    summary.sort(key=lambda x: x["page"])

    elapsed = time.perf_counter() - stage_start
    log_translation_summary(summary, elapsed)

    return summary


def log_translation_summary(summary, elapsed: float):
    success_count = sum(1 for x in summary if x.get("status") == "success")
    skipped_count = sum(1 for x in summary if x.get("status") == "skipped")
    failed_count = sum(1 for x in summary if x.get("status") == "failed")

    successful_durations = [
        x.get("duration_seconds", 0)
        for x in summary
        if x.get("status") == "success" and x.get("duration_seconds", 0) > 0
    ]

    if successful_durations:
        avg_page_time = sum(successful_durations) / len(successful_durations)
        pages_per_hour_at_actual_runtime = (len(successful_durations) / elapsed) * 3600 if elapsed > 0 else 0
        log(
            f"Translation stage completed | duration={format_duration(elapsed)} | "
            f"success={success_count} | skipped={skipped_count} | failed={failed_count} | "
            f"avg_success_page_time={format_duration(avg_page_time)} | "
            f"actual_success_pages_per_hour={pages_per_hour_at_actual_runtime:.2f}"
        )
    else:
        log(
            f"Translation stage completed | duration={format_duration(elapsed)} | "
            f"success={success_count} | skipped={skipped_count} | failed={failed_count}"
        )


# ---------------------------------------------------------------------
# Step 3: Extract and inject images
# Based on notebook 03.
# ---------------------------------------------------------------------

def extract_embedded_images(pdf_path: Path, out_dir: Path):
    """
    Extract only images that are actually displayed/rendered on the page.

    Why not page.get_images(full=True)?
    Some PDFs expose a shared/global image resource list to every page.
    In those cases, get_images(full=True) can return hundreds of image
    resources even when the page visually contains few or no images.

    page.get_image_info(xrefs=True) gives rendered image occurrences with
    bounding boxes, which is safer for POP PDFs.
    """
    doc = fitz.open(pdf_path)
    page = doc[0]

    extracted = []
    seen_xrefs = set()

    out_dir.mkdir(parents=True, exist_ok=True)

    try:
        image_infos = page.get_image_info(xrefs=True)
    except Exception as e:
        log(f"{pdf_path.parent.name} | image extraction warning | get_image_info failed | error={e}")
        image_infos = []

    for info in image_infos:
        xref = info.get("xref")

        if not xref:
            continue

        if xref in seen_xrefs:
            continue

        seen_xrefs.add(xref)

        bbox = info.get("bbox")

        if bbox:
            rect = fitz.Rect(bbox)

            # Ignore tiny artifacts, icons, noise, or invisible fragments.
            if rect.width < 30 or rect.height < 30:
                continue

        try:
            img_info = doc.extract_image(xref)
            img_bytes = img_info["image"]
            ext = img_info.get("ext", "png")

            out_path = out_dir / f"image_{len(extracted) + 1}.{ext}"
            out_path.write_bytes(img_bytes)

            extracted.append({
                "index": len(extracted) + 1,
                "xref": xref,
                "path": out_path,
                "bbox": list(bbox) if bbox else None,
            })

        except Exception as e:
            log(f"{pdf_path.parent.name} | image extraction warning | xref={xref} | error={e}")

    doc.close()
    return extracted


def inject_images_into_html(html_path: Path, embedded_images, final_html_path: Path):
    """
    Primary behavior follows notebook 03:
    - find <figure class="image-placeholder">
    - insert image tag at the start
    - remove image-placeholder class
    """
    html_text = html_path.read_text(encoding="utf-8")
    soup = BeautifulSoup(html_text, "html.parser")

    placeholders = soup.find_all("figure", class_="image-placeholder")
    placeholder_count = len(placeholders)
    image_count = len(embedded_images)

    for idx, fig in enumerate(placeholders, start=1):
        if idx > image_count:
            continue

        img_meta = embedded_images[idx - 1]
        img_rel_path = f"images/{img_meta['path'].name}"

        for content in list(fig.contents):
            if isinstance(content, NavigableString) and "[IMAGE]" in str(content):
                content.extract()

        img_tag = soup.new_tag("img")
        img_tag["src"] = img_rel_path
        img_tag["alt"] = f"Image {idx}"
        img_tag["style"] = "max-width: 100%; height: auto; display: block; margin: 0 auto;"
        fig.insert(0, img_tag)

        existing_classes = fig.get("class", [])
        fig["class"] = [c for c in existing_classes if c != "image-placeholder"]

    # Safe fallback:
    # If Gemini returns plain [IMAGE] instead of <figure class="image-placeholder">,
    # replace plain placeholders in order. This does not affect notebook-style outputs.
    if placeholder_count == 0 and embedded_images:
        html_text_after_primary = str(soup)

        replaced = 0
        for idx, img_meta in enumerate(embedded_images, start=1):
            if "[IMAGE]" not in html_text_after_primary:
                break

            img_rel_path = f"images/{img_meta['path'].name}"
            img_block = (
                f'<figure>'
                f'<img src="{img_rel_path}" alt="Image {idx}" '
                f'style="max-width: 100%; height: auto; display: block; margin: 0 auto;" />'
                f'</figure>'
            )

            html_text_after_primary = html_text_after_primary.replace("[IMAGE]", img_block, 1)
            replaced += 1

        if replaced > 0:
            soup = BeautifulSoup(html_text_after_primary, "html.parser")
            placeholder_count = replaced

    final_html_path.write_text(str(soup), encoding="utf-8")

    return {
        "placeholder_count": placeholder_count,
        "embedded_image_count": image_count,
        "matched_count": min(placeholder_count, image_count),
    }


def html_has_image_placeholders(html_text: str) -> bool:
    """
    Return True only if the translated HTML explicitly expects images.

    This prevents storage blowups for PDFs where every split page exposes
    hundreds of shared image resources even when the page does not actually
    need image injection.
    """
    if "[IMAGE]" in html_text:
        return True

    soup = BeautifulSoup(html_text, "html.parser")

    if soup.find("figure", class_="image-placeholder"):
        return True

    return False


def process_image_injection_for_page(
    page_pdf_path: Path,
    overwrite_existing: bool,
):
    page_start = time.perf_counter()

    page_dir = page_pdf_path.parent
    pname = page_dir.name

    html_path = page_dir / "translated.html"
    images_dir = page_dir / "images"
    final_html_path = page_dir / "final_with_images.html"
    err_path = page_dir / "image_injection_error.txt"

    if not page_pdf_path.exists():
        msg = f"PDF not found: {page_pdf_path}"
        log(f"{pname} | image injection failed | {msg}")
        return {"page": pname, "status": "failed", "error": msg}

    if not html_path.exists():
        msg = f"translated.html not found: {html_path}"
        log(f"{pname} | image injection failed | {msg}")
        return {"page": pname, "status": "failed", "error": msg}

    if final_html_path.exists() and not overwrite_existing:
        log(f"{pname} | image injection skipped | final_with_images.html already exists")
        return {
            "page": pname,
            "status": "skipped",
            "reason": "final_with_images.html already exists",
            "output_file": str(final_html_path),
            "duration_seconds": 0,
        }

    log(f"{pname} | image injection started")

    try:
        if images_dir.exists() and overwrite_existing:
            shutil.rmtree(images_dir)

        html_text = html_path.read_text(encoding="utf-8")

        # Critical storage fix:
        # If Gemini did not place image placeholders, do not extract PDF images.
        if not html_has_image_placeholders(html_text):
            if images_dir.exists():
                shutil.rmtree(images_dir)

            final_html_path.write_text(html_text, encoding="utf-8")

            elapsed = time.perf_counter() - page_start

            log(
                f"{pname} | image injection skipped internally | "
                f"no image placeholders found | duration={format_duration(elapsed)} | "
                f"saved={final_html_path}"
            )

            return {
                "page": pname,
                "status": "success",
                "embedded_image_count": 0,
                "placeholder_count": 0,
                "matched_count": 0,
                "output_file": str(final_html_path),
                "duration_seconds": round(elapsed, 2),
            }

        embedded_images = extract_embedded_images(page_pdf_path, images_dir)

        injection_result = inject_images_into_html(
            html_path=html_path,
            embedded_images=embedded_images,
            final_html_path=final_html_path,
        )

        elapsed = time.perf_counter() - page_start

        log(
            f"{pname} | image injection finished | "
            f"duration={format_duration(elapsed)} | "
            f"images={len(embedded_images)} | "
            f"placeholders={injection_result['placeholder_count']} | "
            f"matched={injection_result['matched_count']} | "
            f"saved={final_html_path}"
        )

        return {
            "page": pname,
            "status": "success",
            "embedded_image_count": len(embedded_images),
            "placeholder_count": injection_result["placeholder_count"],
            "matched_count": injection_result["matched_count"],
            "output_file": str(final_html_path),
            "duration_seconds": round(elapsed, 2),
        }

    except Exception as e:
        elapsed = time.perf_counter() - page_start
        err = str(e)
        err_path.write_text(err, encoding="utf-8")

        log(
            f"{pname} | image injection failed | "
            f"duration={format_duration(elapsed)} | error={err}"
        )

        return {
            "page": pname,
            "status": "failed",
            "error": err,
            "error_file": str(err_path),
            "duration_seconds": round(elapsed, 2),
        }


def inject_images_for_pages(page_pdf_files, overwrite_existing: bool):
    stage_start = time.perf_counter()
    summary = []

    log(f"Image injection stage started | pages={len(page_pdf_files)}")

    for idx, pdf_path in enumerate(page_pdf_files, start=1):
        log(f"Image injection progress | starting {idx}/{len(page_pdf_files)} | page={pdf_path.parent.name}")

        result = process_image_injection_for_page(
            page_pdf_path=pdf_path,
            overwrite_existing=overwrite_existing,
        )

        summary.append(result)

        log(
            f"Image injection progress | completed {idx}/{len(page_pdf_files)} | "
            f"page={pdf_path.parent.name} | status={result.get('status')}"
        )

    elapsed = time.perf_counter() - stage_start

    success_count = sum(1 for x in summary if x.get("status") == "success")
    skipped_count = sum(1 for x in summary if x.get("status") == "skipped")
    failed_count = sum(1 for x in summary if x.get("status") == "failed")

    log(
        f"Image injection stage completed | duration={format_duration(elapsed)} | "
        f"success={success_count} | skipped={skipped_count} | failed={failed_count}"
    )

    return summary


# ---------------------------------------------------------------------
# Step 4: Build combined HTML for DOCX
# Uses styling inspired by notebook 04 print HTML.
# ---------------------------------------------------------------------

def build_combined_docx_html(
    page_pdf_files,
    workdir_root: Path,
    doc_title: str,
    start_page: int,
    end_page: int,
):
    stage_start = time.perf_counter()
    sections = []

    log(f"Combined HTML build started | pages={len(page_pdf_files)}")

    for idx, pdf_path in enumerate(page_pdf_files, start=1):
        page_dir = pdf_path.parent
        final_html_path = page_dir / "final_with_images.html"

        if not final_html_path.exists():
            raise FileNotFoundError(f"Missing final_with_images.html: {final_html_path}")

        fragment = final_html_path.read_text(encoding="utf-8")
        soup = BeautifulSoup(fragment, "html.parser")

        # Rewrite image src paths from page-local path to combined-html-relative path.
        for img in soup.find_all("img"):
            src = img.get("src")
            if not src:
                continue

            src_path = Path(src)

            if src_path.is_absolute():
                abs_img_path = src_path
            else:
                abs_img_path = (page_dir / src_path).resolve()

            rel_img_path = os.path.relpath(abs_img_path, start=workdir_root.resolve())
            img["src"] = rel_img_path.replace("\\", "/")

        content = soup.body.decode_contents() if soup.body else str(soup)

        sections.append(
            f"""
<section class="translated-page" id="{page_dir.name}">
{content}
</section>
"""
        )

        if idx % 25 == 0 or idx == len(page_pdf_files):
            log(f"Combined HTML build progress | added {idx}/{len(page_pdf_files)} pages")

    combined_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>{doc_title}</title>
  <style>
    body {{
      font-family:
        "Noto Sans",
        "Noto Sans Devanagari",
        "DejaVu Sans",
        "Liberation Sans",
        Arial,
        sans-serif;
      font-size: 11pt;
      line-height: 1.4;
      color: #111;
    }}

    h1, h2, h3, h4, h5, h6 {{
      margin-top: 14px;
      margin-bottom: 8px;
    }}

    p {{
      margin: 0 0 10px 0;
    }}

    table {{
      border-collapse: collapse;
      width: 100%;
      table-layout: fixed;
      margin: 8px 0 12px 0;
    }}

    th, td {{
      border: 1px solid #444;
      padding: 6px;
      vertical-align: top;
      word-wrap: break-word;
      overflow-wrap: anywhere;
    }}

    figure {{
      margin: 14px 0;
      text-align: center;
      page-break-inside: avoid;
      break-inside: avoid;
    }}

    img {{
      max-width: 100%;
      height: auto;
      display: block;
      margin: 0 auto;
    }}

    figcaption {{
      margin-top: 8px;
      font-size: 10pt;
    }}

    ul, ol {{
      margin-top: 6px;
      margin-bottom: 10px;
    }}

    .translated-page {{
      page-break-after: always;
      break-after: page;
    }}

    .translated-page:last-child {{
      page-break-after: auto;
      break-after: auto;
    }}
  </style>
</head>
<body>
{chr(10).join(sections)}
</body>
</html>
"""

    combined_html_path = workdir_root / f"{safe_docx_name(doc_title)}_combined_pages_{start_page:03d}_to_{end_page:03d}.html"
    combined_html_path.write_text(combined_html, encoding="utf-8")

    elapsed = time.perf_counter() - stage_start

    log(
        f"Combined HTML build completed | duration={format_duration(elapsed)} | "
        f"saved={combined_html_path}"
    )

    return combined_html_path


# ---------------------------------------------------------------------
# Step 5: Convert combined HTML to DOCX with Pandoc
# ---------------------------------------------------------------------

def check_pandoc_available():
    try:
        subprocess.run(
            ["pandoc", "--version"],
            check=True,
            capture_output=True,
            text=True,
        )
    except Exception as e:
        raise RuntimeError(
            "Pandoc is required for HTML → DOCX conversion, but it was not found.\n"
            "Install Pandoc and make sure 'pandoc' works in CMD/PowerShell.\n"
            "Check with: pandoc --version"
        ) from e


def convert_html_to_docx(
    combined_html_path: Path,
    output_docx_path: Path,
    reference_docx: Path | None = None,
    check_availability: bool = True,
):
    stage_start = time.perf_counter()

    if check_availability:
        check_pandoc_available()

    cmd = [
        "pandoc",
        str(combined_html_path.name),
        "-f",
        "html",
        "-t",
        "docx",
        "-o",
        str(output_docx_path.resolve()),
    ]

    if reference_docx:
        cmd.insert(-2, f"--reference-doc={str(reference_docx.resolve())}")

    log("[RUN] " + " ".join(cmd))

    subprocess.run(
        cmd,
        cwd=str(combined_html_path.parent),
        check=True,
    )

    elapsed = time.perf_counter() - stage_start

    log(
        f"DOCX conversion completed | duration={format_duration(elapsed)} | "
        f"saved={output_docx_path}"
    )

    return output_docx_path


# ---------------------------------------------------------------------
# Step 4/5 Alternative: page-wise DOCX conversion + final merge
# ---------------------------------------------------------------------

def build_single_page_docx_html(
    page_pdf_path: Path,
    doc_title: str,
):
    """
    Build a standalone HTML file for one translated page.

    The HTML is written inside the page folder so relative image paths such as
    images/image_1.png continue to resolve correctly during Pandoc conversion.
    """
    page_dir = page_pdf_path.parent
    pname = page_dir.name
    final_html_path = page_dir / "final_with_images.html"

    if not final_html_path.exists():
        raise FileNotFoundError(f"Missing final_with_images.html: {final_html_path}")

    fragment = final_html_path.read_text(encoding="utf-8")
    soup = BeautifulSoup(fragment, "html.parser")

    # Keep only visible body content if Gemini returned a complete HTML document.
    content = soup.body.decode_contents() if soup.body else str(soup)

    single_page_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>{doc_title} - {pname}</title>
  <style>
    body {{
      font-family:
        "Noto Sans",
        "Noto Sans Devanagari",
        "DejaVu Sans",
        "Liberation Sans",
        Arial,
        sans-serif;
      font-size: 11pt;
      line-height: 1.4;
      color: #111;
    }}

    h1, h2, h3, h4, h5, h6 {{
      margin-top: 14px;
      margin-bottom: 8px;
    }}

    p {{
      margin: 0 0 10px 0;
    }}

    table {{
      border-collapse: collapse;
      width: 100%;
      table-layout: fixed;
      margin: 8px 0 12px 0;
    }}

    th, td {{
      border: 1px solid #444;
      padding: 6px;
      vertical-align: top;
      word-wrap: break-word;
      overflow-wrap: anywhere;
    }}

    figure {{
      margin: 14px 0;
      text-align: center;
      page-break-inside: avoid;
      break-inside: avoid;
    }}

    img {{
      max-width: 100%;
      height: auto;
      display: block;
      margin: 0 auto;
    }}

    figcaption {{
      margin-top: 8px;
      font-size: 10pt;
    }}

    ul, ol {{
      margin-top: 6px;
      margin-bottom: 10px;
    }}

    .translated-page {{
      page-break-after: always;
      break-after: page;
    }}
  </style>
</head>
<body>
<section class="translated-page" id="{pname}">
{content}
</section>
</body>
</html>
"""

    single_page_html_path = page_dir / f"{pname}_single_page_docx.html"
    single_page_html_path.write_text(single_page_html, encoding="utf-8")
    return single_page_html_path


def convert_page_to_docx(
    page_pdf_path: Path,
    output_docx_path: Path,
    doc_title: str,
    overwrite_existing: bool,
    reference_docx: Path | None = None,
):
    page_start = time.perf_counter()
    page_dir = page_pdf_path.parent
    pname = page_dir.name

    if output_docx_path.exists() and not overwrite_existing:
        log(f"{pname} | page DOCX skipped | already exists: {output_docx_path}")
        return {
            "page": pname,
            "status": "skipped",
            "output_file": str(output_docx_path),
            "duration_seconds": 0,
        }

    try:
        log(f"{pname} | page DOCX build started")

        single_page_html_path = build_single_page_docx_html(
            page_pdf_path=page_pdf_path,
            doc_title=doc_title,
        )

        convert_html_to_docx(
            combined_html_path=single_page_html_path,
            output_docx_path=output_docx_path,
            reference_docx=reference_docx,
            check_availability=False,
        )

        elapsed = time.perf_counter() - page_start
        log(
            f"{pname} | page DOCX build finished | "
            f"duration={format_duration(elapsed)} | saved={output_docx_path}"
        )

        return {
            "page": pname,
            "status": "success",
            "html_file": str(single_page_html_path),
            "output_file": str(output_docx_path),
            "duration_seconds": round(elapsed, 2),
        }

    except Exception as e:
        elapsed = time.perf_counter() - page_start
        err_path = page_dir / "page_docx_error.txt"
        err_path.write_text(str(e), encoding="utf-8")

        log(
            f"{pname} | page DOCX build failed | "
            f"duration={format_duration(elapsed)} | error={e}"
        )

        return {
            "page": pname,
            "status": "failed",
            "error": str(e),
            "error_file": str(err_path),
            "duration_seconds": round(elapsed, 2),
        }


def convert_pages_to_individual_docx(
    page_pdf_files,
    final_output_dir: Path,
    safe_name: str,
    doc_title: str,
    overwrite_existing: bool,
    reference_docx: Path | None = None,
):
    """
    Convert every page's final_with_images.html into its own DOCX file.

    This avoids building one huge combined HTML before conversion. It is more
    robust for large POP PDFs and for pages whose HTML can cause Pandoc to
    truncate a combined DOCX.
    """
    stage_start = time.perf_counter()
    page_docx_dir = final_output_dir / "page_docx"
    page_docx_dir.mkdir(parents=True, exist_ok=True)

    log(f"Page-wise DOCX stage started | pages={len(page_pdf_files)} | output_dir={page_docx_dir}")

    check_pandoc_available()

    results = []
    page_docx_files = []

    for idx, pdf_path in enumerate(page_pdf_files, start=1):
        pname = pdf_path.parent.name
        page_docx_path = page_docx_dir / f"{safe_name}_{pname}.docx"

        log(f"Page-wise DOCX progress | starting {idx}/{len(page_pdf_files)} | page={pname}")

        result = convert_page_to_docx(
            page_pdf_path=pdf_path,
            output_docx_path=page_docx_path,
            doc_title=doc_title,
            overwrite_existing=overwrite_existing,
            reference_docx=reference_docx,
        )

        results.append(result)

        if result.get("status") in {"success", "skipped"}:
            page_docx_files.append(page_docx_path)

        log(
            f"Page-wise DOCX progress | completed {idx}/{len(page_pdf_files)} | "
            f"page={pname} | status={result.get('status')}"
        )

    elapsed = time.perf_counter() - stage_start

    success_count = sum(1 for x in results if x.get("status") == "success")
    skipped_count = sum(1 for x in results if x.get("status") == "skipped")
    failed_count = sum(1 for x in results if x.get("status") == "failed")

    log(
        f"Page-wise DOCX stage completed | duration={format_duration(elapsed)} | "
        f"success={success_count} | skipped={skipped_count} | failed={failed_count}"
    )

    return page_docx_files, results


def merge_docx_files_in_order(input_docx_files, output_docx_path: Path):
    """Merge page-wise DOCX files into one final DOCX in exact list order."""
    stage_start = time.perf_counter()

    input_docx_files = [Path(p) for p in input_docx_files]

    missing = [str(p) for p in input_docx_files if not p.exists()]
    if missing:
        raise FileNotFoundError("Missing page DOCX files:\n" + "\n".join(missing))

    if not input_docx_files:
        raise ValueError("No page DOCX files were provided for merging.")

    output_docx_path.parent.mkdir(parents=True, exist_ok=True)

    log(f"DOCX merge started | files={len(input_docx_files)}")
    log(f"DOCX merge first file: {input_docx_files[0]}")
    log(f"DOCX merge last file: {input_docx_files[-1]}")

    master = Document(str(input_docx_files[0]))
    composer = Composer(master)

    for idx, docx_path in enumerate(input_docx_files[1:], start=2):
        composer.append(Document(str(docx_path)))

        if idx % 25 == 0 or idx == len(input_docx_files):
            log(f"DOCX merge progress | appended {idx}/{len(input_docx_files)} files")

    composer.save(str(output_docx_path))

    elapsed = time.perf_counter() - stage_start
    log(
        f"DOCX merge completed | duration={format_duration(elapsed)} | "
        f"saved={output_docx_path}"
    )

    return output_docx_path


# ---------------------------------------------------------------------
# Main CLI
# ---------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="POP Translation Pipeline: source PDF -> Gemini 3.1 Pro HTML -> images injected -> page-wise DOCX -> merged final DOCX"
    )

    parser.add_argument("--source-pdf", required=True, help="Path to source POP PDF")
    parser.add_argument("--workdir-root", required=True, help="Working/output directory for this document")
    parser.add_argument("--doc-name", default=None, help="Document name for final output")
    parser.add_argument("--prompt-file", default="prompts/page_to_pdf.txt", help="Prompt file path")
    parser.add_argument("--model", default="gemini-3.1-pro-preview", help="Gemini model name")

    parser.add_argument("--start-page", type=int, default=1, help="Start page, 1-indexed")
    parser.add_argument("--end-page", type=int, default=None, help="End page, 1-indexed. Default: last page")

    parser.add_argument("--overwrite", action="store_true", help="Overwrite existing outputs")
    parser.add_argument("--max-retries", type=int, default=5, help="Gemini max retries per page")
    parser.add_argument("--retry-wait-seconds", type=int, default=15, help="Retry wait seconds")

    parser.add_argument("--concurrency", type=int, default=1, help="Parallel translation workers. Use 1 for notebook-equivalent behavior.")
    parser.add_argument("--disable-google-search", action="store_true", help="Disable Google Search tool. Default keeps notebook behavior enabled.")

    parser.add_argument("--reference-docx", default=None, help="Optional Pandoc reference DOCX")
    parser.add_argument("--skip-translation", action="store_true", help="Skip Gemini translation and reuse existing translated.html files")
    parser.add_argument("--skip-image-injection", action="store_true", help="Skip image injection and reuse existing final_with_images.html files")

    args = parser.parse_args()
    pipeline_start = time.perf_counter()

    source_pdf_path = Path(args.source_pdf).resolve()
    workdir_root = Path(args.workdir_root).resolve()
    prompt_path = Path(args.prompt_file).resolve() if args.prompt_file else None
    reference_docx = Path(args.reference_docx).resolve() if args.reference_docx else None

    doc_name = args.doc_name or source_pdf_path.stem
    safe_name = safe_docx_name(doc_name)

    workdir_root.mkdir(parents=True, exist_ok=True)
    set_log_file(workdir_root / "pipeline_runtime.log")

    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key and not args.skip_translation:
        raise RuntimeError("GEMINI_API_KEY is not set.")

    print("=" * 90)
    print("POP Translation Pipeline")
    print("=" * 90)
    print("Source PDF        :", source_pdf_path)
    print("Workdir root      :", workdir_root)
    print("Doc name          :", doc_name)
    print("Prompt file       :", prompt_path)
    print("Model             :", args.model)
    print("Thinking level    : HIGH")
    print("Google Search     :", not args.disable_google_search)
    print("Start page        :", args.start_page)
    print("End page          :", args.end_page if args.end_page else "LAST PAGE")
    print("Concurrency       :", args.concurrency)
    print("Overwrite         :", args.overwrite)
    print("Runtime log       :", workdir_root / "pipeline_runtime.log")
    print("=" * 90)

    log("Pipeline configuration loaded")
    log(f"Source PDF: {source_pdf_path}")
    log(f"Workdir root: {workdir_root}")
    log(f"Doc name: {doc_name}")
    log(f"Model: {args.model}")
    log("Thinking level: HIGH")
    log(f"Google Search: {not args.disable_google_search}")
    log(f"Start page: {args.start_page}")
    log(f"End page: {args.end_page if args.end_page else 'LAST PAGE'}")
    log(f"Concurrency: {args.concurrency}")
    log(f"Overwrite: {args.overwrite}")

    # 1. Split PDF
    log("[1/5] Splitting PDF into page-wise PDFs...")

    page_pdf_files, total_pages = split_pdf_to_page_folders(
        source_pdf_path=source_pdf_path,
        output_root=workdir_root,
        start_page=args.start_page,
        end_page=args.end_page,
        overwrite_existing=args.overwrite,
    )

    actual_end_page = args.end_page if args.end_page else total_pages

    log(f"Created or found {len(page_pdf_files)} page PDFs")
    log("First 5 page PDFs:")
    for p in page_pdf_files[:5]:
        log(f"  {p}")

    log("Last 5 page PDFs:")
    for p in page_pdf_files[-5:]:
        log(f"  {p}")

    # 2. Translate PDF pages
    translation_summary = []

    if args.skip_translation:
        log("[2/5] Skipping translation. Reusing existing translated.html files.")
    else:
        log("[2/5] Translating pages to HTML using notebook-equivalent Gemini config...")
        prompt = load_prompt(prompt_path)

        translation_summary = translate_pages(
            page_pdf_files=page_pdf_files,
            prompt=prompt,
            model=args.model,
            api_key=api_key,
            overwrite_existing=args.overwrite,
            max_retries=args.max_retries,
            retry_wait_seconds=args.retry_wait_seconds,
            enable_google_search=not args.disable_google_search,
            concurrency=args.concurrency,
        )

        translation_summary_path = workdir_root / "translation_summary.json"
        translation_summary_path.write_text(
            json.dumps(translation_summary, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        log(f"Translation summary saved: {translation_summary_path}")

        failed = [x for x in translation_summary if x["status"] == "failed"]
        if failed:
            raise RuntimeError(f"Translation failed for {len(failed)} page(s). See translation_summary.json")

    # 3. Extract/inject images
    image_summary = []

    if args.skip_image_injection:
        log("[3/5] Skipping image injection. Reusing existing final_with_images.html files.")
    else:
        log("[3/5] Extracting and injecting images...")
        image_summary = inject_images_for_pages(
            page_pdf_files=page_pdf_files,
            overwrite_existing=args.overwrite,
        )

        image_summary_path = workdir_root / "image_injection_summary.json"
        image_summary_path.write_text(
            json.dumps(image_summary, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        log(f"Image injection summary saved: {image_summary_path}")

        failed = [x for x in image_summary if x["status"] == "failed"]
        if failed:
            raise RuntimeError(f"Image injection failed for {len(failed)} page(s). See image_injection_summary.json")

    # 4. Convert each page HTML to one page-wise DOCX
    log("[4/5] Converting each page HTML to page-wise DOCX...")

    final_output_dir = workdir_root / "final_output"
    final_output_dir.mkdir(parents=True, exist_ok=True)

    page_docx_files, page_docx_summary = convert_pages_to_individual_docx(
        page_pdf_files=page_pdf_files,
        final_output_dir=final_output_dir,
        safe_name=safe_name,
        doc_title=doc_name,
        overwrite_existing=args.overwrite,
        reference_docx=reference_docx,
    )

    page_docx_summary_path = workdir_root / "page_docx_summary.json"
    page_docx_summary_path.write_text(
        json.dumps(page_docx_summary, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    log(f"Page-wise DOCX summary saved: {page_docx_summary_path}")

    failed = [x for x in page_docx_summary if x.get("status") == "failed"]
    if failed:
        raise RuntimeError(f"Page-wise DOCX conversion failed for {len(failed)} page(s). See page_docx_summary.json")

    # 5. Merge page-wise DOCX files into final DOCX
    log("[5/5] Merging page-wise DOCX files into final DOCX...")

    final_docx_path = final_output_dir / f"{safe_name}_translated_pages_{args.start_page:03d}_to_{actual_end_page:03d}.docx"

    merge_docx_files_in_order(
        input_docx_files=page_docx_files,
        output_docx_path=final_docx_path,
    )

    total_elapsed = time.perf_counter() - pipeline_start

    log("Pipeline completed successfully")
    log(f"Final DOCX: {final_docx_path}")
    log(f"Total pipeline time: {format_duration(total_elapsed)}")

    print("\nDONE")
    print("Final DOCX:", final_docx_path)
    print("Total pipeline time:", format_duration(total_elapsed))
    print("Runtime log:", workdir_root / "pipeline_runtime.log")


if __name__ == "__main__":
    main()